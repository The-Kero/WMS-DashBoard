# -*- coding: utf-8 -*-
from openpyxl import load_workbook

wb = load_workbook(r'C:\Users\JWPark\Desktop\냉장 재고조사\251207 냉장 재고조사(신선식품).xlsx')
ws = wb.active

print('='*60)
print('내가 만든 파일: 251207 냉장 재고조사(신선식품).xlsx')
print('='*60)

print('\n=== 1행 셀별 상세 ===')
for col in ['A','B','C','D','E','F','G','H']:
    c = ws[f'{col}1']
    print(f'{col}1: value="{c.value}"')
    print(f'     font: {c.font.name}, {c.font.size}pt, bold={c.font.bold}')
    if c.fill.patternType:
        if c.fill.fgColor.rgb:
            print(f'     fill: pattern={c.fill.patternType}, rgb={c.fill.fgColor.rgb}')
        else:
            print(f'     fill: pattern={c.fill.patternType}, theme={c.fill.fgColor.theme}')
    else:
        print(f'     fill: None')
    print(f'     align: h={c.alignment.horizontal}, v={c.alignment.vertical}')
    print()

print('\n=== 2행 셀별 상세 ===')
for col in ['A','B','C','D','E','F','G','H']:
    c = ws[f'{col}2']
    print(f'{col}2: value="{c.value}"')
    print(f'     font: {c.font.name}, {c.font.size}pt, bold={c.font.bold}')
    if c.fill.patternType:
        if c.fill.fgColor.rgb:
            print(f'     fill: pattern={c.fill.patternType}, rgb={c.fill.fgColor.rgb}')
        else:
            print(f'     fill: pattern={c.fill.patternType}, theme={c.fill.fgColor.theme}')
    else:
        print(f'     fill: None')
    print()

print('\n=== 3행 (데이터행) 셀별 상세 ===')
for col in ['A','B','C','D','E','F','G','H']:
    c = ws[f'{col}3']
    print(f'{col}3: value="{c.value}"')
    print(f'     font: {c.font.name}, {c.font.size}pt, bold={c.font.bold}')
    if c.fill.patternType:
        print(f'     fill: pattern={c.fill.patternType}')
    else:
        print(f'     fill: None')
    print()

print('\n=== 인쇄 설정 ===')
print(f'반복행: {ws.print_title_rows}')
print(f'머리글 우측: {ws.oddHeader.right.text}')
print(f'바닥글 우측: {ws.oddFooter.right.text}')
