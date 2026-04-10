# -*- coding: utf-8 -*-
"""
재고조사 엑셀 자동 생성 모듈
- 초록불 조건 충족 시 부진재고/신선식품 엑셀 파일 생성
- v39 설계서 기반 구현
- 작성일: 2025-12-07
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.worksheet.page import PageMargins
from openpyxl.utils.dataframe import dataframe_to_rows
from datetime import datetime
from pathlib import Path
import os

# ========================================
# 설정값
# ========================================
INVENTORY_CSV_DIR = r"C:\OSIS_AUTO\inventory_status"
FILTER_DIR = r"C:\Users\JWPark\Desktop\냉장 재고조사"
OUTPUT_DIR = r"C:\Users\JWPark\Desktop\냉장 재고조사"

# ========================================
# 데이터 로드 함수
# ========================================
def load_inventory_csv(date_str=None):
    """재고현황 CSV 로드"""
    if date_str is None:
        date_str = datetime.now().strftime("%Y%m%d")
    
    file_path = Path(INVENTORY_CSV_DIR) / f"inventory_status_{date_str}.csv"
    
    if not file_path.exists():
        raise FileNotFoundError(f"재고현황 파일 없음: {file_path}")
    
    df = pd.read_csv(file_path, encoding='utf-8-sig')
    return df

def load_filter_csv(filter_type):
    """필터 CSV 로드 (신선식품 또는 부진재고) - 헤더 없음"""
    if filter_type == "fresh":
        file_path = Path(FILTER_DIR) / "신선식품.csv"
        col_name = "상품"
    elif filter_type == "slow":
        file_path = Path(FILTER_DIR) / "부진재고.csv"
        col_name = "로케이션"
    else:
        raise ValueError(f"알 수 없는 필터 타입: {filter_type}")
    
    if not file_path.exists():
        raise FileNotFoundError(f"필터 파일 없음: {file_path}")
    
    # 헤더 없는 CSV → 컬럼명 지정
    df = pd.read_csv(file_path, encoding='utf-8-sig', header=None, names=[col_name])
    return df

# ========================================
# 필터 및 그룹핑 함수
# ========================================
def filter_fresh_products(inventory_df, fresh_filter_df):
    """신선식품 필터링: 상품코드 기준"""
    fresh_codes = fresh_filter_df['상품'].astype(str).tolist()
    filtered = inventory_df[inventory_df['상품'].astype(str).isin(fresh_codes)].copy()
    return filtered

def filter_slow_moving(inventory_df, slow_filter_df, fresh_filter_df):
    """부진재고 필터링: 로케이션 기준 + 신선식품 제외"""
    slow_locations = slow_filter_df['로케이션'].astype(str).tolist()
    fresh_codes = fresh_filter_df['상품'].astype(str).tolist()
    
    # 로케이션 필터 적용
    filtered = inventory_df[inventory_df['로케이션'].astype(str).isin(slow_locations)].copy()
    # 신선식품 제외
    filtered = filtered[~filtered['상품'].astype(str).isin(fresh_codes)]
    return filtered

def group_fresh_products(df):
    """신선식품 그룹핑 및 합산"""
    grouped = df.groupby(
        ['로케이션', '상품', '상품명', '소비기한'],
        as_index=False
    ).agg({'가용수량': 'sum'})
    return grouped

def group_slow_moving(df):
    """부진재고 그룹핑 및 합산"""
    grouped = df.groupby(
        ['로케이션', '상품', '상품명', '단위및규격', '소비기한', '입수량'],
        as_index=False
    ).agg({
        '가용수량': 'sum',
        '가용박스수량': 'sum',
        '가용잔량': 'sum'
    })
    return grouped



# ========================================
# 엑셀 스타일 설정
# ========================================
def get_styles():
    """공통 스타일 정의 - 12개 차이점 반영 완료"""
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    # 헤더용 (10pt bold)
    header1_font = Font(name='맑은 고딕', size=10, bold=True)
    # 헤더2행용 (9pt bold)
    header2_font = Font(name='맑은 고딕', size=9, bold=True)
    # 데이터용 (9pt)
    data_font = Font(name='맑은 고딕', size=9)
    # 데이터용 bold (F열, G열용)
    data_bold_font = Font(name='맑은 고딕', size=9, bold=True)
    # 병합셀 내부용 (굴림체 9pt)
    merge_inner_font = Font(name='굴림', size=9)
    # G1 차이 헤더용 (9pt bold - 기존 10pt에서 수정)
    g1_font = Font(name='맑은 고딕', size=9, bold=True)
    # H2 비고 부제목용 (10pt bold - 기존 9pt에서 수정)
    h2_font = Font(name='맑은 고딕', size=10, bold=True)
    
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')
    
    # 헤더 배경색 (올리브 계열 - theme 3 + tint -0.25 근사값)
    header_fill = PatternFill(start_color='C4D79B', end_color='C4D79B', fill_type='solid')
    # F열 데이터 배경색 (연한 파랑 - theme 4 + tint 0.80 근사값)
    f_col_fill = PatternFill(start_color='DCE6F1', end_color='DCE6F1', fill_type='solid')
    
    return {
        'border': thin_border,
        'header1_font': header1_font,
        'header2_font': header2_font,
        'data_font': data_font,
        'data_bold_font': data_bold_font,
        'merge_inner_font': merge_inner_font,
        'g1_font': g1_font,
        'h2_font': h2_font,
        'center': center_align,
        'left': left_align,
        'header_fill': header_fill,
        'f_col_fill': f_col_fill
    }

# ========================================
# 부진재고 엑셀 생성 (12열)
# ========================================
def create_slow_moving_excel(df, output_path):
    """부진재고 엑셀 파일 생성 - 12개 차이점 반영 완료"""
    wb = Workbook()
    ws = wb.active
    ws.title = "냉장"
    styles = get_styles()
    
    # 열 너비 설정 (샘플 파일 정확값)
    col_widths = {
        'A': 6.5, 'B': 15.6640625, 'C': 18.83203125, 'D': 47.0, 'E': 18.33203125,
        'F': 15.1640625, 'G': 8.5, 'H': 12.83203125, 'I': 13.0, 'J': 13.0,
        'K': 13.0, 'L': 26.6640625
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width
    
    # 행 높이 설정
    ws.row_dimensions[1].height = 20.25
    ws.row_dimensions[2].height = 20.25
    
    # [핵심] 병합 전에 모든 헤더 셀에 스타일 먼저 적용
    for row in range(1, 3):
        for col in range(1, 13):
            cell = ws.cell(row=row, column=col)
            if row == 1:
                cell.font = styles['header1_font']  # 10pt bold
            else:
                cell.font = styles['header2_font']  # 9pt bold
            cell.alignment = styles['center']
            cell.border = styles['border']
            cell.fill = styles['header_fill']  # 올리브색 배경
    
    # [수정3] 병합셀 내부에 굴림체 9pt 적용 (병합 전!)
    for cell_addr in ['C1', 'D1', 'E1', 'G1', 'H1', 'I1', 'J1', 'A2', 'K2']:
        ws[cell_addr].font = styles['merge_inner_font']
    
    # [수정1] K1 폰트 크기 9pt로 변경 (G1과 동일 역할)
    ws['K1'].font = styles['g1_font']
    
    # [수정2] L2 폰트 크기 10pt로 변경 (H2와 동일 역할)
    ws['L2'].font = styles['h2_font']
    
    # 헤더 1행 - 병합셀 (스타일 적용 후 병합!)
    ws.merge_cells('A1:A2')  # No.
    ws.merge_cells('B1:E1')  # 제품정보
    ws.merge_cells('F1:J1')  # 전산재고
    ws.merge_cells('K1:K2')  # 차이
    
    # 헤더 값 입력
    ws['A1'] = 'No.'
    ws['B1'] = '제품정보'
    ws['F1'] = '전산재고'
    ws['K1'] = '차이'
    ws['L1'] = '비고'
    
    # 헤더 2행 값 입력
    headers_row2 = {2: '로케이션', 3: '상품', 4: '상품명', 5: '단위및규격', 
                   6: '소비기한', 7: '입수량', 8: '가용수량', 9: 'BOX', 10: '낱개', 12: '(특이사항 기재)'}
    for col_idx, header in headers_row2.items():
        ws.cell(row=2, column=col_idx, value=header)

    # 데이터 입력 (3행부터)
    for idx, row_data in enumerate(df.itertuples(), start=1):
        row_num = idx + 2  # 3행부터 시작
        ws.row_dimensions[row_num].height = 22.5
        
        # No. (수식)
        if idx == 1:
            ws.cell(row=row_num, column=1, value=1)
        else:
            ws.cell(row=row_num, column=1, value=f'=A{row_num-1}+1')
        
        # B열: 로케이션
        cell_b = ws.cell(row=row_num, column=2, value=row_data.로케이션)
        cell_b.number_format = '0_);[Red]\\(0\\)'
        
        # [수정4] C열: 상품 - 문자열로 저장
        cell_c = ws.cell(row=row_num, column=3, value=str(row_data.상품))
        
        # D열: 상품명
        cell_d = ws.cell(row=row_num, column=4, value=row_data.상품명)
        cell_d.number_format = '#,##0_);(#,##0)'
        
        # E열: 단위및규격
        ws.cell(row=row_num, column=5, value=row_data.단위및규격)
        
        # [수정5] F열: 소비기한 - datetime 변환
        exp_str = str(int(row_data.소비기한))
        try:
            exp_date = datetime.strptime(exp_str, "%Y%m%d")
            cell_f = ws.cell(row=row_num, column=6, value=exp_date)
        except:
            cell_f = ws.cell(row=row_num, column=6, value=exp_str)
        cell_f.number_format = 'mm-dd-yy'
        
        # G열: 입수량
        cell_g = ws.cell(row=row_num, column=7, value=row_data.입수량)
        cell_g.number_format = '#,##0_);(#,##0)'
        
        # [수정6] H열: 가용수량 - bold + 배경색 + 천단위
        cell_h = ws.cell(row=row_num, column=8, value=row_data.가용수량)
        cell_h.font = styles['data_bold_font']
        cell_h.fill = styles['f_col_fill']
        cell_h.number_format = '#,##0_);(#,##0)'
        
        # I열: BOX (가용박스수량)
        cell_i = ws.cell(row=row_num, column=9, value=row_data.가용박스수량)
        cell_i.number_format = '#,##0_);(#,##0)'
        
        # J열: 낱개 (가용잔량)
        cell_j = ws.cell(row=row_num, column=10, value=row_data.가용잔량)
        cell_j.number_format = '#,##0_);(#,##0)'
        
        # [수정7] K열: 차이 - bold + 텍스트 서식
        cell_k = ws.cell(row=row_num, column=11, value='')
        cell_k.font = styles['data_bold_font']
        cell_k.number_format = '@'
        
        # [수정9] L열: 비고 - 텍스트 서식
        cell_l = ws.cell(row=row_num, column=12, value='')
        cell_l.number_format = '@'
        
        # 스타일 적용
        for col in range(1, 13):
            cell = ws.cell(row=row_num, column=col)
            if col not in [8, 11]:  # H, K는 이미 bold 적용됨
                cell.font = styles['data_font']
            cell.border = styles['border']
            # [수정8,9] 모든 컬럼 center 정렬
            cell.alignment = styles['center']
    
    # [수정12] 인쇄 설정
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = 9  # A4
    ws.page_margins = PageMargins(
        top=0.5905511811023623,
        bottom=0.5905511811023623,
        left=0.7874015748031497,
        right=0.3937007874015748
    )
    ws.print_title_rows = '1:2'  # 반복 인쇄 행
    ws.oddHeader.right.text = "재고조사일 : &D, &T"
    ws.oddFooter.right.text = "&P  /  &N"
    
    # 저장
    wb.save(output_path)
    return output_path


# ========================================
# 신선식품 엑셀 생성 (8열)
# ========================================
def create_fresh_products_excel(df, output_path):
    """신선식품 엑셀 파일 생성 - 12개 차이점 반영 완료"""
    wb = Workbook()
    ws = wb.active
    ws.title = "냉장"
    styles = get_styles()
    
    # 열 너비 설정 (샘플 파일 정확값)
    col_widths = {
        'A': 6.5, 'B': 15.6640625, 'C': 18.83203125, 'D': 47.0,
        'E': 15.1640625, 'F': 12.83203125, 'G': 13.0, 'H': 64.1640625
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width
    
    # 행 높이 설정
    ws.row_dimensions[1].height = 20.25
    ws.row_dimensions[2].height = 20.25
    
    # [핵심] 병합 전에 모든 헤더 셀에 스타일 먼저 적용
    for row in range(1, 3):
        for col in range(1, 9):
            cell = ws.cell(row=row, column=col)
            if row == 1:
                cell.font = styles['header1_font']  # 10pt bold
            else:
                cell.font = styles['header2_font']  # 9pt bold
            cell.alignment = styles['center']
            cell.border = styles['border']
            cell.fill = styles['header_fill']  # 올리브색 배경
    
    # [수정3] 병합셀 내부에 굴림체 9pt 적용 (병합 전!)
    for cell_addr in ['C1', 'D1', 'F1', 'A2', 'G2']:
        ws[cell_addr].font = styles['merge_inner_font']
    
    # [수정1] G1 폰트 크기 9pt로 변경
    ws['G1'].font = styles['g1_font']
    
    # [수정2] H2 폰트 크기 10pt로 변경
    ws['H2'].font = styles['h2_font']
    
    # 헤더 1행 - 병합셀 (스타일 적용 후 병합!)
    ws.merge_cells('A1:A2')  # No.
    ws.merge_cells('B1:D1')  # 제품정보
    ws.merge_cells('E1:F1')  # 전산재고
    ws.merge_cells('G1:G2')  # 차이
    
    # 헤더 값 입력
    ws['A1'] = 'No.'
    ws['B1'] = '제품정보'
    ws['E1'] = '전산재고'
    ws['G1'] = '차이'
    ws['H1'] = '비고'
    
    # 헤더 2행 값 입력
    headers_row2 = {2: '로케이션', 3: '상품', 4: '상품명', 5: '소비기한', 6: '가용수량', 8: '(특이사항 기재)'}
    for col_idx, header in headers_row2.items():
        ws.cell(row=2, column=col_idx, value=header)

    # 데이터 입력 (3행부터)
    for idx, row_data in enumerate(df.itertuples(), start=1):
        row_num = idx + 2  # 3행부터 시작
        ws.row_dimensions[row_num].height = 21.75
        
        # No. (수식)
        if idx == 1:
            ws.cell(row=row_num, column=1, value=1)
        else:
            ws.cell(row=row_num, column=1, value=f'=A{row_num-1}+1')
        
        # B열: 로케이션
        cell_b = ws.cell(row=row_num, column=2, value=row_data.로케이션)
        cell_b.number_format = '0_);[Red]\\(0\\)'
        
        # [수정4] C열: 상품 - 문자열로 저장
        cell_c = ws.cell(row=row_num, column=3, value=str(row_data.상품))
        
        # D열: 상품명
        cell_d = ws.cell(row=row_num, column=4, value=row_data.상품명)
        cell_d.number_format = '#,##0_);(#,##0)'
        
        # [수정5] E열: 소비기한 - datetime 변환
        exp_str = str(int(row_data.소비기한))
        try:
            exp_date = datetime.strptime(exp_str, "%Y%m%d")
            cell_e = ws.cell(row=row_num, column=5, value=exp_date)
        except:
            cell_e = ws.cell(row=row_num, column=5, value=exp_str)
        cell_e.number_format = 'mm-dd-yy'
        
        # [수정6] F열: 가용수량 - bold + 배경색 + 천단위
        cell_f = ws.cell(row=row_num, column=6, value=row_data.가용수량)
        cell_f.font = styles['data_bold_font']
        cell_f.fill = styles['f_col_fill']
        cell_f.number_format = '#,##0_);(#,##0)'
        
        # [수정7] G열: 차이 - bold + 텍스트 서식
        cell_g = ws.cell(row=row_num, column=7, value='')
        cell_g.font = styles['data_bold_font']
        cell_g.number_format = '@'
        
        # [수정9] H열: 비고 - 텍스트 서식
        cell_h = ws.cell(row=row_num, column=8, value='')
        cell_h.number_format = '@'
        
        # 스타일 적용
        for col in range(1, 9):
            cell = ws.cell(row=row_num, column=col)
            if col not in [6, 7]:  # F, G는 이미 bold 적용됨
                cell.font = styles['data_font']
            cell.border = styles['border']
            # [수정8,9] 모든 컬럼 center 정렬
            cell.alignment = styles['center']
    
    # [수정12] 인쇄 설정
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = 9  # A4
    ws.page_margins = PageMargins(
        top=0.5905511811023623,
        bottom=0.5905511811023623,
        left=0.7874015748031497,
        right=0.3937007874015748
    )
    ws.print_title_rows = '1:2'  # 반복 인쇄 행
    ws.oddHeader.right.text = "재고조사일 : &D, &T"
    ws.oddFooter.right.text = "&P  /  &N"
    
    # 저장
    wb.save(output_path)
    return output_path


# ========================================
# 메인 실행 함수
# ========================================
def generate_inventory_excel(date_str=None):
    """
    재고조사 엑셀 파일 생성 메인 함수
    - 부진재고, 신선식품 2개 파일 생성
    - 반환: {'success': bool, 'files': list, 'message': str}
    """
    if date_str is None:
        date_str = datetime.now().strftime("%Y%m%d")
    
    # 파일명 형식: YYMMDD
    file_date = date_str[2:]  # 20251207 → 251207
    
    result = {
        'success': False,
        'files': [],
        'message': ''
    }
    
    try:
        # 1. 데이터 로드
        print(f"[1/5] 재고현황 CSV 로드 중... ({date_str})")
        inventory_df = load_inventory_csv(date_str)
        print(f"      → {len(inventory_df)}행 로드 완료")
        
        # 2. 필터 파일 로드
        print("[2/5] 필터 파일 로드 중...")
        fresh_filter = load_filter_csv("fresh")
        slow_filter = load_filter_csv("slow")
        print(f"      → 신선식품 {len(fresh_filter)}개, 부진재고 {len(slow_filter)}개")
        
        # 3. 신선식품 처리
        print("[3/5] 신선식품 엑셀 생성 중...")
        fresh_data = filter_fresh_products(inventory_df, fresh_filter)
        fresh_grouped = group_fresh_products(fresh_data)
        fresh_sorted = fresh_grouped.sort_values(['로케이션', '상품', '소비기한'])
        
        fresh_filename = f"{file_date} 냉장 재고조사(신선식품).xlsx"
        fresh_path = Path(OUTPUT_DIR) / fresh_filename
        create_fresh_products_excel(fresh_sorted, fresh_path)
        result['files'].append(str(fresh_path))
        print(f"      → {fresh_filename} 생성 완료 ({len(fresh_sorted)}행)")
        
        # 4. 부진재고 처리
        print("[4/5] 부진재고 엑셀 생성 중...")
        slow_data = filter_slow_moving(inventory_df, slow_filter, fresh_filter)
        slow_grouped = group_slow_moving(slow_data)
        slow_sorted = slow_grouped.sort_values(['로케이션', '상품', '소비기한'])
        
        slow_filename = f"{file_date} 냉장 재고조사(부진재고).xlsx"
        slow_path = Path(OUTPUT_DIR) / slow_filename
        create_slow_moving_excel(slow_sorted, slow_path)
        result['files'].append(str(slow_path))
        print(f"      → {slow_filename} 생성 완료 ({len(slow_sorted)}행)")
        
        # 5. 완료
        print("[5/5] 완료!")
        result['success'] = True
        result['message'] = f"재고조사 엑셀 2개 파일 생성 완료"
        
    except FileNotFoundError as e:
        result['message'] = f"파일 없음: {e}"
        print(f"[오류] {result['message']}")
    except Exception as e:
        result['message'] = f"생성 실패: {e}"
        print(f"[오류] {result['message']}")
    
    return result


# ========================================
# 독립 실행
# ========================================
if __name__ == '__main__':
    print("=" * 50)
    print("재고조사 엑셀 자동 생성 시작")
    print("=" * 50)
    
    result = generate_inventory_excel()
    
    print()
    print("=" * 50)
    print("결과:")
    print(f"  성공 여부: {result['success']}")
    print(f"  메시지: {result['message']}")
    if result['files']:
        print("  생성된 파일:")
        for f in result['files']:
            print(f"    - {f}")
    print("=" * 50)
