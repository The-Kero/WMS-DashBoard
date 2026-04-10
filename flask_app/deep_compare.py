# -*- coding: utf-8 -*-
from openpyxl import load_workbook

print('='*70)
print('DEEP COMPARE: OLD vs NEW')
print('='*70)

wb_old = load_workbook(r'C:\Users\JWPark\Desktop\냉장 재고조사\251205 냉장 재고조사(신선식품).xlsx')
wb_new = load_workbook(r'C:\Users\JWPark\Desktop\냉장 재고조사\251207 냉장 재고조사(신선식품).xlsx')
ws_old = wb_old.active
ws_new = wb_new.active

print('\n[1] BASIC INFO')
print(f'Sheet: OLD={ws_old.title} / NEW={ws_new.title}')
print(f'Range: OLD={ws_old.dimensions} / NEW={ws_new.dimensions}')

print('\n[2] MERGED CELLS')
old_merged = sorted([str(m) for m in ws_old.merged_cells.ranges])
new_merged = sorted([str(m) for m in ws_new.merged_cells.ranges])
print(f'OLD: {old_merged}')
print(f'NEW: {new_merged}')

print('\n[3] COLUMN WIDTH')
for col in ['A','B','C','D','E','F','G','H']:
    w_old = ws_old.column_dimensions[col].width
    w_new = ws_new.column_dimensions[col].width
    status = 'OK' if abs(w_old - w_new) < 0.01 else 'DIFF'
    print(f'{col}: OLD={w_old:.2f} / NEW={w_new:.2f} [{status}]')

print('\n[4] ROW HEIGHT (1-10)')
for row in range(1, 11):
    h_old = ws_old.row_dimensions[row].height
    h_new = ws_new.row_dimensions[row].height
    status = 'OK' if h_old == h_new else 'DIFF'
    print(f'Row{row}: OLD={h_old} / NEW={h_new} [{status}]')


print('\n[5] BORDER STYLE (Row 1-5)')
for row in range(1, 6):
    for col in ['A','B','C','D','E','F','G','H']:
        c_old = ws_old[f'{col}{row}']
        c_new = ws_new[f'{col}{row}']
        b_old = c_old.border
        b_new = c_new.border
        
        old_sides = (b_old.left.style if b_old.left else None,
                     b_old.right.style if b_old.right else None,
                     b_old.top.style if b_old.top else None,
                     b_old.bottom.style if b_old.bottom else None)
        new_sides = (b_new.left.style if b_new.left else None,
                     b_new.right.style if b_new.right else None,
                     b_new.top.style if b_new.top else None,
                     b_new.bottom.style if b_new.bottom else None)
        
        if old_sides != new_sides:
            print(f'{col}{row}: OLD={old_sides} / NEW={new_sides} [DIFF]')

print('\n[6] ALIGNMENT (Row 1-5)')
for row in range(1, 6):
    for col in ['A','B','C','D','E','F','G','H']:
        c_old = ws_old[f'{col}{row}']
        c_new = ws_new[f'{col}{row}']
        
        old_align = (c_old.alignment.horizontal, c_old.alignment.vertical, c_old.alignment.wrap_text)
        new_align = (c_new.alignment.horizontal, c_new.alignment.vertical, c_new.alignment.wrap_text)
        
        if old_align != new_align:
            print(f'{col}{row}: OLD={old_align} / NEW={new_align} [DIFF]')

print('\n[7] FILL/BACKGROUND (Row 1-5)')
for row in range(1, 6):
    for col in ['A','B','C','D','E','F','G','H']:
        c_old = ws_old[f'{col}{row}']
        c_new = ws_new[f'{col}{row}']
        
        if c_old.fill.patternType:
            if c_old.fill.fgColor.theme is not None:
                old_fill = f"theme={c_old.fill.fgColor.theme},tint={c_old.fill.fgColor.tint:.2f}"
            else:
                old_fill = f"rgb={c_old.fill.fgColor.rgb}"
        else:
            old_fill = 'None'
            
        if c_new.fill.patternType:
            if c_new.fill.fgColor.theme is not None:
                new_fill = f"theme={c_new.fill.fgColor.theme},tint={c_new.fill.fgColor.tint:.2f}"
            else:
                new_fill = f"rgb={c_new.fill.fgColor.rgb}"
        else:
            new_fill = 'None'
        
        if old_fill != new_fill:
            print(f'{col}{row}: OLD={old_fill} / NEW={new_fill} [DIFF]')


print('\n[8] FONT DETAIL (Row 1-5)')
for row in range(1, 6):
    print(f'--- Row {row} ---')
    for col in ['A','B','C','D','E','F','G','H']:
        c_old = ws_old[f'{col}{row}']
        c_new = ws_new[f'{col}{row}']
        
        old_font = f"{c_old.font.name},{c_old.font.size}pt,bold={c_old.font.bold}"
        new_font = f"{c_new.font.name},{c_new.font.size}pt,bold={c_new.font.bold}"
        
        if old_font != new_font:
            print(f'  {col}: OLD=[{old_font}] / NEW=[{new_font}] [DIFF]')

print('\n[9] NUMBER FORMAT (Row 3-7)')
for row in range(3, 8):
    for col in ['A','B','C','D','E','F','G','H']:
        c_old = ws_old[f'{col}{row}']
        c_new = ws_new[f'{col}{row}']
        
        if c_old.number_format != c_new.number_format:
            print(f'{col}{row}: OLD="{c_old.number_format}" / NEW="{c_new.number_format}" [DIFF]')

print('\n[10] VALUE TYPE (Row 3-5)')
for row in range(3, 6):
    for col in ['A','B','C','D','E','F','G','H']:
        c_old = ws_old[f'{col}{row}']
        c_new = ws_new[f'{col}{row}']
        
        old_type = type(c_old.value).__name__
        new_type = type(c_new.value).__name__
        
        if old_type != new_type:
            print(f'{col}{row}: OLD={old_type}({c_old.value}) / NEW={new_type}({c_new.value}) [DIFF]')

print('\n[11] PRINT SETTINGS')
print(f'Title Rows: OLD={ws_old.print_title_rows} / NEW={ws_new.print_title_rows}')
print(f'Orientation: OLD={ws_old.page_setup.orientation} / NEW={ws_new.page_setup.orientation}')
print(f'Paper Size: OLD={ws_old.page_setup.paperSize} / NEW={ws_new.page_setup.paperSize}')
print(f'Margin Top: OLD={ws_old.page_margins.top} / NEW={ws_new.page_margins.top}')
print(f'Margin Bottom: OLD={ws_old.page_margins.bottom} / NEW={ws_new.page_margins.bottom}')
print(f'Margin Left: OLD={ws_old.page_margins.left} / NEW={ws_new.page_margins.left}')
print(f'Margin Right: OLD={ws_old.page_margins.right} / NEW={ws_new.page_margins.right}')

print('\n[12] HEADER/FOOTER')
print(f'Header Left: OLD="{ws_old.oddHeader.left.text}" / NEW="{ws_new.oddHeader.left.text}"')
print(f'Header Center: OLD="{ws_old.oddHeader.center.text}" / NEW="{ws_new.oddHeader.center.text}"')
print(f'Header Right: OLD="{ws_old.oddHeader.right.text}" / NEW="{ws_new.oddHeader.right.text}"')
print(f'Footer Left: OLD="{ws_old.oddFooter.left.text}" / NEW="{ws_new.oddFooter.left.text}"')
print(f'Footer Center: OLD="{ws_old.oddFooter.center.text}" / NEW="{ws_new.oddFooter.center.text}"')
print(f'Footer Right: OLD="{ws_old.oddFooter.right.text}" / NEW="{ws_new.oddFooter.right.text}"')

print('\n[13] A COLUMN (No.) - FORMULA CHECK')
for row in range(3, 8):
    c_old = ws_old[f'A{row}']
    c_new = ws_new[f'A{row}']
    print(f'A{row}: OLD={c_old.value}(type={type(c_old.value).__name__}) / NEW={c_new.value}(type={type(c_new.value).__name__})')

print('\n[14] DATA ROW 3 FULL COMPARE')
print('OLD Row 3:')
for col in ['A','B','C','D','E','F','G','H']:
    c = ws_old[f'{col}3']
    print(f'  {col}: {c.value}')
print('NEW Row 3:')
for col in ['A','B','C','D','E','F','G','H']:
    c = ws_new[f'{col}3']
    print(f'  {col}: {c.value}')

print('\n' + '='*70)
print('DONE')
print('='*70)
